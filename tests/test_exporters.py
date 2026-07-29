"""
Pytest Unit Tests for Master Excel (.xlsx) and CSV Index Exporter Engine.
"""
from __future__ import annotations

import csv
import tempfile
from pathlib import Path

from prc_checker.exporters import HAS_OPENPYXL, generate_excel_and_csv
from prc_checker.models import CardResult


def test_excel_and_csv_export():
    with tempfile.TemporaryDirectory() as tmpdir:
        outdir = Path(tmpdir)
        results = [
            CardResult(cts="1640", village="Byculla", district="23", status="ok", file_jpg="propcard_Byculla_1640.jpg", bytes=407102),
            CardResult(cts="1641", village="Byculla", district="23", status="ok", file_jpg="propcard_Byculla_1641.jpg", bytes=443941)
        ]
        generate_excel_and_csv(outdir, "Byculla", "23", results)

        # 1. Verify CSV Export
        csv_file = outdir / "index.csv"
        assert csv_file.exists()
        with open(csv_file, "r", encoding="utf-8") as f:
            reader = list(csv.DictReader(f))
            assert len(reader) == 2
            assert reader[0]["cts"] == "1640"
            assert reader[0]["status"] == "ok"
            assert reader[1]["cts"] == "1641"

        # 2. Verify Excel Workbook (.xlsx) Export
        excel_file = outdir / "index.xlsx"
        assert excel_file.exists()
        if HAS_OPENPYXL:
            import openpyxl
            wb = openpyxl.load_workbook(excel_file)
            ws = wb["Query Index"]
            assert ws.cell(row=2, column=1).value == "1640"
            assert ws.cell(row=2, column=2).value == "OK"
            assert ws.cell(row=3, column=1).value == "1641"
            wb.close()
