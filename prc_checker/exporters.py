"""
Data Export Engines: Master Excel Index (.xlsx) and CSV Summary Exporter.
Accepts list[CardResult] Pydantic models.
"""
from __future__ import annotations

import csv
from pathlib import Path

from .models import CardResult

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


def generate_excel_and_csv(outdir: Path, village: str, district: str, results: list[CardResult]):
    """Export batch execution results to index.xlsx and index.csv master query index."""
    # 1. Export CSV Master Index & Summary Alias
    fieldnames = [
        "cts", "status", "village", "district",
        "file_jpg", "file_html", "file_report",
        "file_promoter_brief", "file_lawyer_brief", "file_json",
        "bytes", "detail"
    ]
    
    rows = []
    for r in results:
        data = r.model_dump()
        rows.append({
            "cts": data.get("cts", ""),
            "status": data.get("status", ""),
            "village": village,
            "district": district,
            "file_jpg": data.get("file_jpg", ""),
            "file_html": data.get("file_html", ""),
            "file_report": data.get("file_report", ""),
            "file_promoter_brief": data.get("file_promoter_brief", ""),
            "file_lawyer_brief": data.get("file_lawyer_brief", ""),
            "file_json": data.get("file_json", ""),
            "bytes": data.get("bytes", 0),
            "detail": data.get("detail", "Success")
        })

    for fname in ("index.csv", "summary.csv"):
        csv_path = outdir / fname
        with open(csv_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(rows)

    # 2. Export Master Excel Index (.xlsx)
    excel_path = outdir / "index.xlsx"
    if HAS_OPENPYXL:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Query Index"

        # Define Styles
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        cell_font = Font(name="Segoe UI", size=10)
        link_font = Font(name="Segoe UI", size=10, color="2563EB", underline="single")
        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0")
        )

        headers = [
            "CTS Number", "Status", "Village", "District",
            "Card Image (.jpg)", "HTML Viewer (.html)", "Title Report (.md)",
            "Promoter Brief (.md)", "Lawyer Brief (.md)", "JSON Metadata (.json)",
            "File Size (Bytes)", "Detail / Notes"
        ]
        ws.append(headers)

        for col_num, _ in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center

        for r_idx, row_data in enumerate(rows, 2):
            ws.append([
                row_data["cts"],
                row_data["status"].upper(),
                row_data["village"],
                row_data["district"],
                row_data["file_jpg"],
                row_data["file_html"],
                row_data["file_report"],
                row_data["file_promoter_brief"],
                row_data["file_lawyer_brief"],
                row_data["file_json"],
                row_data["bytes"],
                row_data["detail"]
            ])

            for c_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.font = cell_font
                cell.border = thin_border
                cell.alignment = align_left if c_idx in (12,) else align_center

                # Status Highlighting
                if c_idx == 2:
                    if row_data["status"] == "ok":
                        cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
                        cell.font = Font(name="Segoe UI", size=10, bold=True, color="166534")
                    else:
                        cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                        cell.font = Font(name="Segoe UI", size=10, bold=True, color="991B1B")

                # Hyperlinks for generated files
                if c_idx in (5, 6, 7, 8, 9, 10) and cell.value:
                    filename = str(cell.value)
                    cell.hyperlink = filename
                    cell.font = link_font

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        wb.save(excel_path)


def generate_sqlite(outdir: Path, village: str, district: str, results: list[CardResult]):
    """Export batch execution results to relational SQLite database database.sqlite."""
    import sqlite3
    db_path = outdir / "database.sqlite"
    conn = sqlite3.connect(db_path)
    cur = conn.cursor()
    cur.execute("""
        CREATE TABLE IF NOT EXISTS property_cards (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            cts TEXT NOT NULL,
            village TEXT NOT NULL,
            district TEXT NOT NULL,
            status TEXT NOT NULL,
            file_jpg TEXT,
            file_html TEXT,
            file_report TEXT,
            file_json TEXT,
            bytes INTEGER,
            fetched_at DATETIME DEFAULT CURRENT_TIMESTAMP,
            detail TEXT
        )
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_cts_village ON property_cards(village, cts)")
    for r in results:
        data = r.model_dump()
        cur.execute("""
            INSERT INTO property_cards (cts, village, district, status, file_jpg, file_html, file_report, file_json, bytes, detail)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            str(data.get("cts", "")),
            village,
            district,
            str(data.get("status", "")),
            data.get("file_jpg", ""),
            data.get("file_html", ""),
            data.get("file_report", ""),
            data.get("file_json", ""),
            data.get("bytes", 0),
            data.get("detail", "Success")
        ))
    conn.commit()
    conn.close()

